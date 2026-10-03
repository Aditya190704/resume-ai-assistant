"""Extract clean text from PDF, DOCX and TXT resumes (all in memory)."""

import io
import logging
import os
import re

import pymupdf
from docx import Document

from config import settings

logger = logging.getLogger(__name__)


class ResumeParserError(Exception):
    """Base class for all resume parsing problems (messages are user-friendly)."""


class UnsupportedFileTypeError(ResumeParserError):
    pass


class FileTooLargeError(ResumeParserError):
    pass


class EmptyResumeError(ResumeParserError):
    pass


def normalize_whitespace(text: str) -> str:
    """Remove junk characters and collapse extra spaces / blank lines."""
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t\r\f\v]+", " ", text)
    text = re.sub(r" ?\n ?", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_text_from_pdf(file_bytes: bytes) -> str:
    try:
        with pymupdf.open(stream=file_bytes, filetype="pdf") as doc:
            if doc.needs_pass:
                raise ResumeParserError(
                    "This PDF is password protected. Please upload an unlocked file."
                )

            return "\n".join(page.get_text("text") for page in doc)

    except ResumeParserError:
        raise

    except Exception as exc:
        logger.exception("PDF extraction failed")
        raise ResumeParserError(
            "Unable to read the uploaded PDF."
        ) from exc


def extract_text_from_docx(file_bytes: bytes) -> str:
    try:
        document = Document(io.BytesIO(file_bytes))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        return "\n".join(parts)
    except Exception as exc:
        logger.exception("DOCX extraction failed")
        raise ResumeParserError("Unable to read the uploaded DOCX file.") from exc


def extract_text_from_txt(file_bytes: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return file_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue
    return file_bytes.decode("utf-8", errors="replace")


def extract_resume_text(filename: str, file_bytes: bytes) -> str:
    """Validate the file, extract its text and return clean text."""
    extension = os.path.splitext(filename)[1].lower()

    if extension not in settings.SUPPORTED_EXTENSIONS:
        raise UnsupportedFileTypeError(
            "Unsupported file type. Please upload a PDF, DOCX or TXT file."
        )
    if len(file_bytes) == 0:
        raise EmptyResumeError("The uploaded file is empty.")
    if len(file_bytes) > settings.MAX_FILE_SIZE_BYTES:
        raise FileTooLargeError(
            f"File is too large. Please upload a file smaller than "
            f"{settings.MAX_FILE_SIZE_MB} MB."
        )

    if extension == ".pdf":
        raw_text = extract_text_from_pdf(file_bytes)
    elif extension == ".docx":
        raw_text = extract_text_from_docx(file_bytes)
    else:
        raw_text = extract_text_from_txt(file_bytes)

    text = normalize_whitespace(raw_text)
    if not text:
        raise EmptyResumeError(
            "No readable text found in the resume. If it is a scanned image PDF, "
            "please upload a text-based PDF, DOCX or TXT."
        )
    return text